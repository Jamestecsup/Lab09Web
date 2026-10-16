"""Genera GLAB-S09-...-RESUELTO.docx:
1. Inserta las 6 capturas del alumno en las pruebas Postman (reemplaza las
   2 de ejemplo de la guia por las suyas en p4/p6).
2. Responde las preguntas de los casos 1-6 + observaciones + conclusiones.
"""
import copy
import shutil
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

SRC = "GLAB-S09-RCOELLO-2026-01-1.docx"
DST = "GLAB-S09-RCOELLO-2026-01-1-RESUELTO.docx"
IMG = {
    "p1": "imagenes_postman/p1_signup_admin.png",
    "p2": "imagenes_postman/p2_signup_user.png",
    "p3": "imagenes_postman/p3_signin_admin.png",
    "p4": "imagenes_postman/p4_ruta_admin.png",
    "p5": "imagenes_postman/p5_signin_user.png",
    "p6": "imagenes_postman/p6_ruta_user.png",
}
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

shutil.copyfile(SRC, DST)
doc = Document(DST)
paras = doc.paragraphs
print("total paras:", len(paras))

def txt(p):
    return p.text.strip()

def insert_paragraph_after(ref_para, text="", bold_prefix=None, size=11):
    new_p = copy.deepcopy(ref_para._p)
    # limpiar contenido copiado
    for child in list(new_p):
        new_p.remove(child)
    ref_para._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    np = Paragraph(new_p, ref_para._parent)
    if bold_prefix:
        r = np.add_run(bold_prefix)
        r.bold = True
        r.font.size = Pt(size)
    if text:
        r = np.add_run(text)
        r.font.size = Pt(size)
    return np

def insert_image_after(ref_para, path, width_cm=15, caption=None):
    new_p = copy.deepcopy(ref_para._p)
    for child in list(new_p):
        new_p.remove(child)
    ref_para._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    np = Paragraph(new_p, ref_para._parent)
    np.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    np.add_run().add_picture(path, width=Cm(width_cm))
    if caption:
        cp = insert_paragraph_after(np, caption, size=10)
        cp.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    return np

def clear_drawings(para):
    for d in para._p.findall(f".//{{{W_NS}}}drawing"):
        d.getparent().remove(d)

def replace_image_with(para, path, width_cm=15):
    clear_drawings(para)
    para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    para.add_run().add_picture(path, width=Cm(width_cm))

def find_idx(pred, start=0):
    for i in range(start, len(paras)):
        if pred(txt(paras[i])):
            return i
    return -1

# ---------- 1. IMAGENES POSTMAN ----------
# p1: despues del cierre '}' de la respuesta del signup admin
i_hdr = find_idx(lambda t: t == "1. PROBAR SIGNUP (REGISTRO)")
i_msg = find_idx(lambda t: "Usuario registrado correctamente" in t, i_hdr)
i_close = find_idx(lambda t: t == "}", i_msg)
print("p1:", i_hdr, i_msg, i_close)
insert_image_after(paras[i_close], IMG["p1"], caption="Figura. Registro de admin en Postman (200 OK).")

paras = doc.paragraphs  # refrescar lista
def find_idx2(pred, start=0):
    for i in range(start, len(paras)):
        if pred(paras[i].text.strip()):
            return i
    return -1

# p2: antes del header de la prueba 3 (el test 2 no trae respuesta en la guia)
i_t3 = find_idx2(lambda t: t == "3. PROBAR LOGIN JWT")
print("p2 antes de:", i_t3)
ref = paras[i_t3]
new_p = copy.deepcopy(ref._p)
for child in list(new_p):
    new_p.remove(child)
ref._p.addprevious(new_p)
from docx.text.paragraph import Paragraph
np = Paragraph(new_p, ref._parent)
np.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
np.add_run().add_picture(IMG["p2"], width=Cm(15))
cp2 = insert_paragraph_after(np, "Figura. Registro de user normal en Postman (200 OK).")
cp2.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

paras = doc.paragraphs
# p3: despues del cierre del login admin (tokenType Bearer)
i_t3 = find_idx2(lambda t: t == "3. PROBAR LOGIN JWT")
i_tt = find_idx2(lambda t: '"tokenType": "Bearer"' in t, i_t3)
i_c3 = find_idx2(lambda t: t == "}", i_tt)
print("p3:", i_t3, i_tt, i_c3)
insert_image_after(paras[i_c3], IMG["p3"], caption="Figura. Login admin en Postman (200 OK, accessToken JWT).")

paras = doc.paragraphs
# p4: reemplazar imagen de ejemplo tras GET /api/test/admin
i_get_admin = find_idx2(lambda t: t == "GET: http://localhost:8080/api/test/admin")
i_img4 = -1
for i in range(i_get_admin + 1, min(i_get_admin + 6, len(paras))):
    if paras[i]._p.findall(f".//{{{W_NS}}}drawing"):
        i_img4 = i
        break
print("p4 reemplaza img en:", i_img4)
replace_image_with(paras[i_img4], IMG["p4"])
insert_paragraph_after(paras[i_img4], "Figura. Ruta protegida admin con Bearer Token (200 OK, Admin Content).")

paras = doc.paragraphs
# p5: despues del cierre del login user (segundo tokenType)
i_tu = find_idx2(lambda t: t == "PROBAR USER NORMAL")
occ, i_tt2, count = -1, -1, 0
for i in range(i_tu, len(paras)):
    if '"tokenType": "Bearer"' in paras[i].text:
        count += 1
        if count == 1:
            i_tt2 = i
            break
i_c5 = find_idx2(lambda t: t == "}", i_tt2)
print("p5:", i_tu, i_tt2, i_c5)
insert_image_after(paras[i_c5], IMG["p5"], caption="Figura. Login user normal en Postman (200 OK, ROLE_USER).")

paras = doc.paragraphs
# p6: reemplazar imagen de ejemplo tras GET /api/test/user
i_get_user = find_idx2(lambda t: t == "GET: http://localhost:8080/api/test/user")
i_img6 = -1
for i in range(i_get_user + 1, min(i_get_user + 6, len(paras))):
    if paras[i]._p.findall(f".//{{{W_NS}}}drawing"):
        i_img6 = i
        break
print("p6 reemplaza img en:", i_img6)
replace_image_with(paras[i_img6], IMG["p6"])
insert_paragraph_after(paras[i_img6], "Figura. Ruta protegida user con Bearer Token (200 OK, User Content).")

# ---------- 2. RESPUESTAS ANALISIS ----------
ANS = {
1: [
"Devuelve \"Public Content\" con 200 OK sin necesidad de token.",
"En controller/TestController.java.",
"El metodo allAccess() con @GetMapping(\"/all\").",
"Porque SecurityConfig la declara publica con .requestMatchers(\"/api/test/all\").permitAll().",
"En el metodo filterChain(), dentro de authorizeHttpRequests.",
"La ruta caeria en .anyRequest().authenticated() y sin token devolveria 403 Forbidden.",
"Intercepta cada request y aplica en orden: sin CSRF, sin sesiones (STATELESS), las reglas de autorizacion y el JwtAuthFilter antes del filtro de login.",
],
2: [
"Body vacio con 403 Forbidden.",
"Codigo HTTP 403.",
"La clase jwt/JwtAuthFilter (doFilterInternal llama a parseJwt y no encuentra token).",
"Porque el SecurityContext queda sin Authentication y la regla hasAnyRole(\"USER\",\"ADMIN\") deniega el acceso.",
"No, el rechazo ocurre en la cadena de filtros (AuthorizationFilter) y nunca llega al controller.",
"JwtAuthFilter (extraccion del token) y el filtro de autorizacion (denegacion).",
"Que el header Authorization: Bearer <jwt> esta ausente o el token no valida.",
],
3: [
"Body vacio con 403 Forbidden y en consola sale \"JWT invalido: ...\".",
"El metodo validateJwtToken(String authToken).",
"En la clase jwt/JwtUtils.",
"La libreria io.jsonwebtoken (jjwt-api/impl/jackson) version 0.11.5 del pom.xml.",
"Ocurre una io.jsonwebtoken.JwtException (MalformedJwtException, SignatureException o ExpiredJwtException).",
"Con token inexistente parseJwt devuelve null y ni se valida; con token invalido validateJwtToken devuelve false tras la excepcion. Ambos terminan en 403, pero el invalido deja log en consola.",
"Al fallar no se crea el UsernamePasswordAuthenticationToken ni se guarda nada en el SecurityContextHolder; se sigue la cadena con filterChain.doFilter y la autorizacion posterior deniega con 403.",
],
4: [
"Devuelve \"User Content\" con 200 OK.",
"Se ejecuta userAccess() de TestController.",
"Extrae el Bearer con parseJwt, lo valida con JwtUtils, carga el usuario con UserDetailsServiceImpl, crea el UsernamePasswordAuthenticationToken y lo guarda en el SecurityContextHolder.",
"Desde SecurityContextHolder.getContext().getAuthentication().",
"El objeto UsernamePasswordAuthenticationToken, cuyo principal es UserDetailsImpl y cuyas authorities son ROLE_USER/ROLE_ADMIN.",
"Postman envia Authorization: Bearer <jwt>; JwtAuthFilter extrae el token; JwtUtils lo valida y obtiene el email; UserDetailsServiceImpl carga el usuario; se crea la Authentication; el SecurityContextHolder la guarda; el filtro de autorizacion permite por hasAnyRole; finalmente el controller responde.",
"Si expiro (jwt.expiration = 3600000, 1 hora) salta ExpiredJwtException, la validacion da false y responde 403; hay que hacer signin de nuevo.",
],
5: [
"La API devuelve body vacio con 403 Forbidden.",
"Se obtiene el codigo HTTP 403.",
"En el archivo config/SecurityConfig.java.",
"Con el metodo filterChain() y la regla .requestMatchers(\"/api/test/admin\").hasRole(\"ADMIN\").",
"Del Authentication.getAuthorities() del SecurityContext, lista construida en UserDetailsImpl.build() como \"ROLE_\" + nombre del rol.",
"En una Collection<? extends GrantedAuthority> dentro del UserDetailsImpl y del Authentication.",
"Authentication es quien eres (identidad verificada con un JWT valido) y Authorization es que puedes hacer (roles frente a las reglas). Entre ellas: con token inexistente no hay identidad; con token invalido la identidad falla al validar.",
"Porque su lista de roles es solo [ROLE_USER] y la ruta exige ROLE_ADMIN.",
],
6: [
"El sistema devuelve \"Admin Content\" con 200 OK.",
"Se ejecuta adminAccess() de TestController.",
"Se superaron la firma y vigencia del JWT (JwtUtils) y el rol ADMIN (SecurityConfig).",
"El usuario tenia el rol ADMIN (authority ROLE_ADMIN).",
"Dentro del SecurityContextHolder queda guardado el UsernamePasswordAuthenticationToken con el UserDetailsImpl (id, username, email, authorities) durante toda la request.",
"Signin valida credenciales con el AuthenticationManager (clave con BCrypt); JwtUtils genera el JWT firmado HS256; Postman lo reenvia en el header Authorization; JwtAuthFilter autentica; el filtro de autorizacion permite por rol ADMIN; el controller responde.",
"La clase jwt/JwtUtils con el metodo generateJwtToken (firma HS256 con jwt.secret, subject = email, fechas iat/exp).",
],
}

paras = doc.paragraphs
cur_caso, q_idx = 0, 0
count_ans = 0
i = 0
while i < len(paras):
    t = paras[i].text.strip()
    if t.startswith("Caso 1"):
        cur_caso, q_idx = 1, 0
    elif t.startswith("Caso 2"):
        cur_caso, q_idx = 2, 0
    elif t.startswith("Caso 3"):
        cur_caso, q_idx = 3, 0
    elif t.startswith("Caso 4"):
        cur_caso, q_idx = 4, 0
    elif t.startswith("Caso 5"):
        cur_caso, q_idx = 5, 0
    elif t.startswith("Caso 6"):
        cur_caso, q_idx = 6, 0
    elif t.startswith("OBSERVACIONES") or t.startswith("CONCLUSIONES"):
        cur_caso = 0
    elif cur_caso in ANS and (t.startswith("\u00bf") or t.startswith("Explique el flujo") or t.startswith("Dibuje el flujo")):
        ans = ANS[cur_caso][q_idx]
        np = insert_paragraph_after(paras[i], ans, bold_prefix="Rpta: ")
        count_ans += 1
        q_idx += 1
        paras = doc.paragraphs
        i += 1  # saltar la respuesta insertada
    i += 1
print("respuestas insertadas:", count_ans, "(esperado 7+7+7+7+8+7=43)")

# ---------- 3. OBSERVACIONES Y CONCLUSIONES ----------
OBS = [
"1. Sin MySQL corriendo en XAMPP la aplicacion falla al arrancar con Communications link failure; hay que darle Start a MySQL.",
"2. El application.properties de la guia viene todo pegado en una linea y debe separarse por saltos de linea para que Spring lo lea.",
"3. Sin los roles USER, ADMIN y MODERATOR en la BD el signup falla con Rol no encontrado; se creo un DataInitializer que los genera al arrancar.",
"4. El jwt.secret debe medir 32 o mas caracteres para HS256, si no Keys.hmacShaKeyFor lanza WeakKeyException.",
"5. Pegar el texto <token_admin> literal en vez del JWT real devuelve 403 Forbidden; se debe copiar el accessToken del signin.",
]
CON = [
"1. JWT evita usar sesiones: cada request se valida de forma STATELESS verificando la firma HS256 del token.",
"2. El JwtAuthFilter junto al SecurityContextHolder separan la autenticacion (quien eres) de la autorizacion (que puedes hacer).",
"3. El BCryptPasswordEncoder hace que en la BD nunca se guarden las claves en texto claro.",
"4. Las reglas hasRole y hasAnyRole del SecurityFilterChain protegen las rutas por rol sin codigo extra en los controllers.",
"5. Un token expirado o manipulado se rechaza con 403, por lo que el cliente debe pedir un token nuevo con signin.",
]

paras = doc.paragraphs
def fill_section(starts_with, items):
    plist = doc.paragraphs
    si = -1
    for k in range(len(plist)):
        if plist[k].text.strip().startswith(starts_with):
            si = k
            break
    assert si != -1, starts_with
    j, n = si + 1, 0
    while n < len(items) and j < len(plist):
        tj = plist[j].text.strip()
        if tj.startswith("CONCLUSIONES"):
            break
        if tj and set(tj) <= set("_ "):
            plist[j].text = items[n]
            n += 1
        j += 1
    while n < len(items):
        np = insert_paragraph_after(plist[j - 1], items[n])
        j += 1
        n += 1
        plist = doc.paragraphs
    print(starts_with, "rellenadas:", n)

fill_section("OBSERVACIONES", OBS)
paras = doc.paragraphs
fill_section("CONCLUSIONES", CON)

doc.save(DST)
print("guardado:", DST)
