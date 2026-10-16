"""Pega las 6 capturas de la actividad de analisis al final de cada caso
en GLAB-S09-...-RESUELTO.docx (despues de la ultima Rpta de cada caso).
Requiere: RESUELTO ya generado + imagenes c1..c6 en imagenes_postman/.
"""
import copy
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.text.paragraph import Paragraph

DST = "GLAB-S09-RCOELLO-2026-01-1-RESUELTO.docx"
CASOS = {
    1: ("imagenes_postman/c1_ruta_publica.png",
        "Figura. Caso 1: ruta publica sin token (200 OK, Public Content)."),
    2: ("imagenes_postman/c2_sin_token.png",
        "Figura. Caso 2: acceso sin token (403 Forbidden)."),
    3: ("imagenes_postman/c3_token_invalido.png",
        "Figura. Caso 3: token modificado a mano (403 Forbidden)."),
    4: ("imagenes_postman/c4_token_valido.png",
        "Figura. Caso 4: acceso con token valido de USER (200 OK, User Content)."),
    5: ("imagenes_postman/c5_admin_sin_permiso.png",
        "Figura. Caso 5: token USER en ruta ADMIN (403 Forbidden)."),
    6: ("imagenes_postman/c6_admin_con_permiso.png",
        "Figura. Caso 6: token ADMIN en ruta ADMIN (200 OK, Admin Content)."),
}

doc = Document(DST)
paras = doc.paragraphs

def insert_image_after(ref_para, path, caption):
    new_p = copy.deepcopy(ref_para._p)
    for child in list(new_p):
        new_p.remove(child)
    ref_para._p.addnext(new_p)
    np = Paragraph(new_p, ref_para._parent)
    np.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    np.add_run().add_picture(path, width=Cm(15))
    # pie de figura
    cap_p = copy.deepcopy(new_p)
    for child in list(cap_p):
        cap_p.remove(child)
    new_p.addnext(cap_p)
    cp = Paragraph(cap_p, np._parent)
    cp.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    cp.add_run(caption)

cur, last_rpta = 0, -1
ops = []
for idx, p in enumerate(paras):
    t = p.text.strip()
    if t.startswith("Caso 1:") or t.startswith("Caso 1"):
        if cur and last_rpta != -1:
            ops.append((cur, last_rpta))
        cur, last_rpta = 1, -1
    elif t.startswith("Caso 2"):
        if cur and last_rpta != -1:
            ops.append((cur, last_rpta))
        cur, last_rpta = 2, -1
    elif t.startswith("Caso 3"):
        if cur and last_rpta != -1:
            ops.append((cur, last_rpta))
        cur, last_rpta = 3, -1
    elif t.startswith("Caso 4"):
        if cur and last_rpta != -1:
            ops.append((cur, last_rpta))
        cur, last_rpta = 4, -1
    elif t.startswith("Caso 5"):
        if cur and last_rpta != -1:
            ops.append((cur, last_rpta))
        cur, last_rpta = 5, -1
    elif t.startswith("Caso 6"):
        if cur and last_rpta != -1:
            ops.append((cur, last_rpta))
        cur, last_rpta = 6, -1
    elif t.startswith("OBSERVACIONES"):
        if cur and last_rpta != -1:
            ops.append((cur, last_rpta))
        cur = 0
    elif cur and t.startswith("Rpta:"):
        last_rpta = idx
print("casos detectados:", ops)

# insertar de atras hacia adelante para no mover indices
for caso, idx in sorted(ops, reverse=True):
    path, caption = CASOS[caso]
    insert_image_after(doc.paragraphs[idx], path, caption)
    print(f"caso {caso}: imagen pegada tras parrafo {idx}")

doc.save(DST)
print("guardado:", DST)
